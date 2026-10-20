"""Step 9 acceptance tests: Markdown + Excel exports from reviewed snapshot + manifest."""

from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from portfoliopulse import exports as exports_module
from portfoliopulse import models as m
from portfoliopulse.exports import (
    ExportBlocked,
    ExportWriteError,
    build_markdown_briefing,
    export_run,
)
from portfoliopulse.review import review_claim, review_hypothesis

NOW = datetime(2026, 10, 1, 12, 0, 0)


@pytest.fixture
def approved_run(session):
    """Run through the review cycle: claims approved/edited/rejected + reviewed hypothesis."""
    company = m.Company(
        id=m.new_id("company"),
        name="Nörthstar-Compönnents & Co. (fixture)",
        website="https://northstar.example",
        sector="industrial sensors",
        country="DE",
    )
    session.add(company)
    run = m.ResearchRun(
        id=m.new_id("run"), company_id=company.id, research_focus="integration", status="approved"
    )
    session.add(run)
    source = m.Source(
        id=m.new_id("source"),
        run_id=run.id,
        company_id=company.id,
        url="https://news.example.test/1",
        canonical_url="https://news.example.com/meridian",
        publisher="Trade Press",
        status="fetched",
        content_hash=m.new_id("hash"),
        retrieved_at=NOW,
    )
    session.add(source)
    claim_approved = m.Claim(
        id=m.new_id("claim"),
        run_id=run.id,
        company_id=company.id,
        claim_text="Northstar acquired Meridian for 48 million euros.",
        category="acquisition",
        supporting_excerpt="acquired Meridian for 48 million euros",
        event_date=date(2026, 7, 14),
        verification_status="supported",
        review_status="approved",
    )
    session.add(claim_approved)
    claim_approved.sources.append(source)
    session.flush()
    claim_edited = m.Claim(
        id=m.new_id("claim"),
        run_id=run.id,
        company_id=company.id,
        claim_text="(*) Northstar entered two new regions this year.",
        category="geographic_expansion",
        supporting_excerpt="entered two new regions",
        verification_status="supported",
        review_status="edited",
    )
    session.add(claim_edited)
    claim_edited.sources.append(source)
    claim_rejected = m.Claim(
        id=m.new_id("claim"),
        run_id=run.id,
        company_id=company.id,
        claim_text="The company is secretly restructuring badly.",
        category="strategy",
        supporting_excerpt="excerpt that was never in any source",
        verification_status="unverified",
        review_status="rejected",
    )
    session.add(claim_rejected)
    session.flush()
    hypothesis = m.Hypothesis(
        id=m.new_id("hyp"),
        run_id=run.id,
        company_id=company.id,
        title="Potential integration complexity",
        hypothesis_text="DRAFT — the acquisitions may increase coordination demands.",
        alternative_explanations=["A mature integration function may already exist."],
        missing_evidence=["Operating model and decision rights"],
        questions_to_validate=["Who owns integration decisions?"],
        evidence_summary={"direct_facts": [], "interpretation": {}},
        status="draft",
    )
    session.add(hypothesis)
    hypothesis.claims.append(claim_approved)
    session.commit()
    review_claim(session, claim_id=claim_approved.id, decision="approved")
    review_hypothesis(
        session,
        hypothesis_id=hypothesis.id,
        decision="edited",
        edited_text="The acquisitions may increase coordination demands.",
    )
    return company, run, claim_approved, claim_edited, claim_rejected, hypothesis


# --- gate --------------------------------------------------------------------------


def test_unapproved_runs_cannot_export(session, tmp_path):
    company = m.Company(id=m.new_id("company"), name="Co")
    session.add(company)
    run = m.ResearchRun(
        id=m.new_id("run"), company_id=company.id, research_focus="f", status="awaiting_review"
    )
    session.add(run)
    session.commit()
    with pytest.raises(ExportBlocked, match="not approved"):
        export_run(session, run, out_root=tmp_path)


# --- markdown briefing --------------------------------------------------------------


def test_briefing_contains_the_ten_sections(session, approved_run):
    _company, run, *_rest = approved_run
    text = build_markdown_briefing(session, run)
    for section in (
        "## 1. Company and research scope",
        "## 2. Executive summary",
        "## 3. Company and ownership context",
        "## 4. Material public developments",
        "## 5. Evidence-backed observations",
        "## 6. Preliminary hypotheses to investigate",
        "## 7. Alternative explanations and limitations",
        "## 8. Open questions / recommended next research steps",
        "## 9. Sources and retrieval date",
        "## 10. Review and approval status",
    ):
        assert section in text, section


def test_briefing_uses_reviewer_edits_and_cites_sources(session, approved_run):
    _company, run, claim_approved, claim_edited, _rejected, _hyp = approved_run
    review_claim(
        session,
        claim_id=claim_edited.id,
        decision="edited",
        edited_text="Northstar entered two new regions this year, confirmed.",
    )
    session.commit()
    text = build_markdown_briefing(session, run)
    assert "entered two new regions this year, confirmed." in text  # latest wording
    assert "(*)" not in text
    # established fact with a footnote citation to the source listed in section 9
    assert f"{claim_approved.claim_text} [^" in text
    assert "[Trade Press](https://news.example.com/meridian)" in text


def test_rejected_claims_never_appear_as_facts(session, approved_run):
    _company, run, *_rest = approved_run
    text = build_markdown_briefing(session, run)
    assert "secretly restructuring badly" not in text


def test_hypothesis_sheds_draft_label_once_reviewed(session, approved_run):
    _company, run, *_rest, _hyp = approved_run
    text = build_markdown_briefing(session, run)
    assert "DRAFT —" not in text
    assert "may increase coordination demands." in text


def test_events_without_dates_never_invent_one(session, approved_run):
    _company, run, claim_approved, _edited, _rejected, _hyp = approved_run
    # undated developments ARE listed in section 4, but WITHOUT an invented date;
    # the dated claim carries its real date
    section_4 = build_markdown_briefing(session, run).split("## 4.")[1].split("## 5.")[0]
    dated_bullet = next(
        line
        for line in section_4.splitlines()
        if line.startswith("- ") and "acquired Meridian" in line
    )
    undated_bullet = next(
        line for line in section_4.splitlines() if line.startswith("- ") and "new regions" in line
    )
    assert f"(dated {claim_approved.event_date.isoformat()})" in dated_bullet
    assert "(dated" not in undated_bullet


# --- Excel research pack ------------------------------------------------------------


def test_workbook_sheets_headers_and_labels(session, approved_run, tmp_path):
    company, run, *_rest = approved_run
    del company
    manifest = export_run(session, run, out_root=tmp_path)
    workbook = load_workbook(Path(manifest["directory"]) / "research_pack.xlsx")
    assert workbook.sheetnames == [
        "Executive Summary",
        "Company Profile",
        "Events",
        "Evidence Register",
        "Hypotheses",
        "Open Questions",
    ]
    for name in workbook.sheetnames:
        header_row = [c.value for c in workbook[name][1]]
        assert all(header_row)  # no empty headers
        assert workbook[name].freeze_panes == "A2"
        assert workbook[name].auto_filter.ref  # filters enabled
        for row in workbook[name].iter_rows():
            for cell in row:
                assert "DRAFT —" not in (cell.value or "")


def test_workbook_events_only_dated_and_rejected_absent(session, approved_run, tmp_path):
    _company, run, claim_approved, _edited, _rejected, _hyp = approved_run
    manifest = export_run(session, run, out_root=tmp_path)
    workbook = load_workbook(Path(manifest["directory"]) / "research_pack.xlsx")
    event_rows = list(workbook["Events"].iter_rows(min_row=2, values_only=True))
    assert len(event_rows) == 1  # only the DATED acquisition claim
    assert event_rows[0][0] == claim_approved.event_date.isoformat()
    assert "news.example.com/meridian" in event_rows[0][3]
    register = list(workbook["Evidence Register"].iter_rows(min_row=2, values_only=True))
    assert len(register) == 2  # rejected claim excluded from the register
    assert all("secretly" not in row[2] for row in register)


def test_workbook_urls_present_and_usable_as_links(session, approved_run, tmp_path):
    _company, run, *_rest = approved_run
    manifest = export_run(session, run, out_root=tmp_path)
    workbook = load_workbook(Path(manifest["directory"]) / "research_pack.xlsx")
    evidence = workbook["Evidence Register"]
    url_cells = [
        cell
        for row in evidence.iter_rows(min_row=2)
        for cell in row
        if cell.value and "http" in str(cell.value)
    ]
    assert url_cells
    assert all("news.example.com" in str(cell.value) for cell in url_cells)


# --- manifest, paths, edge cases ----------------------------------------------------


def test_manifest_and_safe_filenames_do_not_collide_across_runs(session, approved_run, tmp_path):
    company, run, *_rest = approved_run
    manifest = export_run(session, run, out_root=tmp_path)
    assert manifest["approval_status"] == "approved"
    assert "analyst_brief.md" in manifest["files"]
    assert "research_pack.xlsx" in manifest["files"]
    assert manifest["generated_at"].endswith("Z")
    expected_dir = tmp_path / f"{run.id[:20]}_n-rthstar-comp-nnents-co-fixture"
    assert (expected_dir / "manifest.json").exists()
    assert (expected_dir / "analyst_brief.md").exists()
    # a second approved run for the same company exports into its own directory
    run_2 = m.ResearchRun(
        id=m.new_id("run"), company_id=company.id, research_focus="expansion", status="approved"
    )
    session.add(run_2)
    session.commit()
    manifest_2 = export_run(session, run_2, out_root=tmp_path)
    assert manifest_2["files"] and manifest_2["run_id"] != manifest["run_id"]


def test_export_failure_recorded_in_manifest_and_raises(
    session, approved_run, tmp_path, monkeypatch
):
    _company, run, *_rest = approved_run
    monkeypatch.setattr(exports_module, "build_workbook", lambda *a, **k: _boom())
    with pytest.raises(ExportWriteError, match="research_pack.xlsx failed"):
        export_run(session, run, out_root=tmp_path)
    manifest_path = tmp_path / f"{run.id[:20]}_n-rthstar-comp-nnents-co-fixture" / "manifest.json"
    import json

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert any("research_pack.xlsx failed" in warning for warning in manifest["warnings"])
    assert "analyst_brief.md" in manifest["files"]  # partial output honestly recorded


def _boom():
    raise RuntimeError("xlsx engine exploded")


def test_nonascii_and_missing_optional_fields(session, approved_run, tmp_path):
    company = m.Company(id=m.new_id("company"), name="Cömpañía Σ	TRACE ÜBER")
    session.add(company)
    run = m.ResearchRun(
        id=m.new_id("run"), company_id=company.id, research_focus="πλάνη", status="approved"
    )
    session.add(run)
    claim = m.Claim(
        id=m.new_id("claim"),
        run_id=run.id,
        company_id=company.id,
        claim_text="Cömpañía está considerando una expansión en Σ europe.",
        category="strategy",
        supporting_excerpt="considerando una expansión",
        event_date=None,
        verification_status="supported",
        review_status="approved",
    )
    session.add(claim)
    session.commit()
    text = build_markdown_briefing(session, run)
    assert "Cömpañía" in text and "Σ" in text  # non-ASCII preserved, not mangled
    manifest = export_run(session, run, out_root=tmp_path)
    assert manifest["files"]
